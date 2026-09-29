"""Extract product references from slides 8-20 of the supplied Showki brochure.

Usage: python scripts/import-showki-brochure.py <brochure.pptx>
The input path is explicit so this importer remains portable.
"""

import io
import json
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'public/assets/products/brochure'
DATA = ROOT / 'src/data/showki-brochure-products.json'
SERIES = {
    8: 'new-arrivals', 9: 'makeup-remover', 10: 'facial-cleansing',
    11: 'deep-cleansing', 12: 'hydration', 13: 'soothing',
    14: 'anti-aging', 15: 'dark-spots', 16: 'mens-skincare',
    17: 'scalp-care', 18: 'body-care', 19: 'sun-protection',
    20: 'foundation',
}
NEW_CATEGORIES = [
    'hydration', 'body-care', 'body-care',
    'deep-cleansing', 'anti-aging', 'sun-protection',
]


def inches(value):
    return value / 914400


def clean(value):
    return re.sub(r'\s+', ' ', value).replace('ﬁ', 'fi').strip()


def main():
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    deck = Presentation(sys.argv[1])
    OUT.mkdir(parents=True, exist_ok=True)
    records = []
    for page in range(8, 21):
        slide = deck.slides[page - 1]
        text = [(shape, clean(shape.text)) for shape in slide.shapes if shape.has_text_frame and clean(shape.text)]
        effects = [(shape, value.split(':', 1)[1].strip()) for shape, value in text if value.lower().startswith('main effects:')]
        specifications = [(shape, value.split(':', 1)[1].strip()) for shape, value in text if value.lower().startswith('specification:')]
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13 and inches(shape.width) < 2.6]
        entries = []
        for effect_shape, effect in effects:
            title_candidates = [(shape, value) for shape, value in text
                                if 0 < inches(effect_shape.top - shape.top) < .8
                                and abs(inches(effect_shape.left - shape.left)) < .6
                                and not value.lower().startswith(('main effects:', 'specification:'))]
            if not title_candidates:
                raise ValueError(f'No title for slide {page}, y={inches(effect_shape.top):.2f}')
            title_shape, name = max(title_candidates, key=lambda item: item[0].top)
            nearest_pictures = sorted(pictures, key=lambda pic: abs(inches(pic.top - title_shape.top)) * 4 + abs(inches(pic.left - title_shape.left) + 1.6))
            if not nearest_pictures:
                raise ValueError(f'No picture for {page}: {name}')
            picture = nearest_pictures[0]
            candidates = [(shape, value) for shape, value in specifications
                          if 0 < inches(shape.top - effect_shape.top) < .8
                          and abs(inches(shape.left - effect_shape.left)) < .6]
            specification = min(candidates, key=lambda item: item[0].top)[1] if candidates else ''
            entries.append((title_shape, name, effect, specification, picture))
        entries.sort(key=lambda item: (round(inches(item[0].top), 1), inches(item[0].left)))
        for number, (_, name, effect, specification, picture) in enumerate(entries, 1):
            image_name = f'p{page:02d}-{number:02d}.webp'
            with Image.open(io.BytesIO(picture.image.blob)) as image:
                image.convert('RGB').save(OUT / image_name, 'WEBP', quality=88, method=6)
            category = NEW_CATEGORIES[number - 1] if page == 8 else SERIES[page]
            records.append({
                'id': f'SK-P{page:02d}-{number:02d}',
                'sourcePage': page,
                'category': category,
                'isNew': page == 8,
                'name': name,
                'effect': effect,
                'specification': specification or None,
                'image': f'/assets/products/brochure/{image_name}',
            })
        print(f'P{page}: {len(entries)} products')
    if len(records) != 57 or len({item['id'] for item in records}) != 57:
        raise ValueError(f'Expected 57 distinct entries; found {len(records)}')
    DATA.write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


if __name__ == '__main__':
    main()
